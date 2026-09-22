import json
import time
from io import BytesIO
from uuid import UUID, uuid4

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

from evalweave.agents.executor import (
    _tabular_markdown_report,
    authenticate_target,
    detect_application_error,
    execute_data_program,
    execute_target_case,
    isolate_session_ids,
    normalize_target_body,
    parse_target_response,
    serialize_results,
    source_parsing_options,
    validate_target,
)
from evalweave.agents.inspection import load_source_rows
from evalweave.agents.model_config import encrypt_secret_payload
from evalweave.agents.planner import evaluate_target_records
from evalweave.agents.python_workspace import run_python_workspace
from evalweave.agents.tools import (
    AssistantToolContext,
    CancelAgentJobTool,
    InspectAgentJobTool,
    RequestOutputFormatTool,
    RunPythonTool,
    SubmitPythonJobTool,
)
from evalweave.agents.workflow import (
    execute_python_job,
    plan_agent_job,
    recover_interrupted_python_jobs,
    repair_python_job_with_react,
    validate_http_target_semantics,
    validate_http_target_with_react,
)
from evalweave.api.routes.agents import (
    infer_explicit_output_format,
    normalize_conversation_title,
    sanitize_assistant_display,
)
from evalweave.core.config import AgentConfig, get_settings
from evalweave.db.models import (
    AgentJob,
    AgentJobStatus,
    AgentStep,
    AssistantConversation,
    AssistantMessage,
    FileObject,
    StepStatus,
)
from evalweave.db.session import get_engine
from evalweave.storage import LocalFileStorage


def test_sanitize_assistant_display_hides_python_workspace_paths() -> None:
    content = (
        "文件在 inputs/ 里，直接复制过去并重命名：\n"
        "文件名已改为 result.xlsx，生成在 outputs/ 目录下"
    )

    rendered = sanitize_assistant_display(content)

    assert "inputs" not in rendered
    assert "outputs" not in rendered
    assert "result.xlsx" in rendered


def test_sanitize_assistant_display_removes_guessed_file_download_url() -> None:
    content = (
        "文件正常了。\n\n"
        "📄 **文件下载：** [结果.xlsx]"
        "(http://example.com/dev/api/project/111/file/222)\n\n"
        "50 条数据均已完成。"
    )

    rendered = sanitize_assistant_display(content)

    assert "文件下载" not in rendered
    assert "http://" not in rendered
    assert "文件正常了" in rendered
    assert "50 条数据均已完成" in rendered


def test_output_format_tool_does_not_ask_again_when_format_is_known() -> None:
    context = AssistantToolContext(
        draft={
            "title": "回答评测",
            "goal": "评测回答质量",
            "source_file_id": str(uuid4()),
            "source_inspected": True,
            "output_format": "xlsx",
        },
        project_id=None,
    )

    result = json.loads(RequestOutputFormatTool().run(context))

    assert result["ok"] is False
    assert context.ui_action is None


@pytest.mark.parametrize(
    ("user_text", "expected"),
    [
        ("结果给我一个 Excel 文件", "xlsx"),
        ("文件名就叫 dev评测结果.xlsx", "xlsx"),
        ("导出 JSONL", "jsonl"),
        ("生成 report.json 文件", "jsonl"),
        ("Markdown 文件就行", "markdown"),
        ("直接在对话中查看，不需要文件", "text"),
        ("调用接口时 contentType 是 TEXT", None),
    ],
)
def test_infer_explicit_output_format(user_text: str, expected: str | None) -> None:
    assert infer_explicit_output_format(user_text) == expected


def login_as_developer(client: TestClient) -> None:
    options = client.get("/api/auth/registration-options").json()["data"]
    development = next(item for item in options if item["code"] == "development")
    client.post(
        "/api/auth/register",
        json={
            "username": "agent_developer",
            "email": "agent_developer@example.com",
            "password": "developer-password",
            "user_type_id": development["id"],
        },
    )
    response = client.post(
        "/api/auth/login",
        json={"username": "agent_developer", "password": "developer-password"},
    )
    assert response.status_code == 200


def current_project(client: TestClient) -> dict[str, object]:
    projects = client.get("/api/projects")
    assert projects.status_code == 200
    return projects.json()["data"][0]


def test_agent_jobs_are_private_to_the_creator(client: TestClient) -> None:
    login_as_developer(client)
    project = current_project(client)
    first_job = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={"title": "First user's task", "goal": "Private evaluation"},
    )
    assert first_job.status_code == 201

    client.post("/api/auth/logout")
    options = client.get("/api/auth/registration-options").json()["data"]
    development = next(item for item in options if item["code"] == "development")
    client.post(
        "/api/auth/register",
        json={
            "username": "other_developer",
            "email": "other_developer@example.com",
            "password": "developer-password",
            "user_type_id": development["id"],
        },
    )
    client.post(
        "/api/auth/login",
        json={"username": "other_developer", "password": "developer-password"},
    )
    second_job = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={"title": "Second user's task", "goal": "Another private evaluation"},
    )
    assert second_job.status_code == 201

    listing = client.get(f"/api/projects/{project['id']}/agent-jobs")
    assert [item["id"] for item in listing.json()["data"]] == [second_job.json()["data"]["id"]]
    first_job_id = first_job.json()["data"]["id"]
    assert client.get(f"/api/agent-jobs/{first_job_id}").status_code == 404
    assert client.get(f"/api/agent-jobs/{first_job_id}/steps").status_code == 404
    assert client.post(f"/api/agent-jobs/{first_job_id}/start").status_code == 404

    client.post("/api/auth/logout")
    client.post(
        "/api/auth/login",
        json={"username": "admin", "password": "test-admin-password"},
    )
    admin_listing = client.get(f"/api/projects/{project['id']}/agent-jobs")
    assert {item["id"] for item in admin_listing.json()["data"]} == {
        first_job_id,
        second_job.json()["data"]["id"],
    }


def test_agent_job_can_be_cancelled_from_api_and_agent_tool(client: TestClient) -> None:
    login_as_developer(client)
    project = current_project(client)
    creator = client.get("/api/auth/me").json()["data"]
    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={"title": "需要停止的任务", "goal": "验证取消能力"},
    ).json()["data"]
    job_id = UUID(created["id"])
    with Session(get_engine()) as session:
        session.add(AgentStep(job_id=job_id, name="execute_eval_spec", status=StepStatus.RUNNING))
        session.add(AgentStep(job_id=job_id, name="summarize", status=StepStatus.PENDING))
        session.commit()

    cancelled = client.post(f"/api/agent-jobs/{job_id}/cancel")

    assert cancelled.status_code == 200
    assert cancelled.json()["data"]["status"] == "cancelled"
    steps = client.get(f"/api/agent-jobs/{job_id}/steps").json()["data"]
    assert [step["status"] for step in steps] == ["cancelled", "cancelled"]

    another = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={"title": "对话中停止", "goal": "由 Agent 停止"},
    ).json()["data"]
    context = AssistantToolContext(
        draft={"background_job_id": another["id"]},
        project_id=UUID(str(project["id"])),
        actor_id=UUID(creator["id"]),
    )

    result = json.loads(CancelAgentJobTool().run(context))

    assert result["ok"] is True
    assert result["status"] == "cancelled"
    assert client.get(f"/api/agent-jobs/{another['id']}").json()["data"]["status"] == "cancelled"


def test_agent_plans_and_executes_without_scheme_approval(client: TestClient) -> None:
    login_as_developer(client)
    runtime = client.get("/api/agent/runtime")
    assert runtime.status_code == 200
    assert "api_key" not in runtime.json()["data"]
    project = current_project(client)
    upload = client.post(
        f"/api/projects/{project['id']}/files",
        files={
            "file": (
                "companion.json",
                b'[{"user_message":"I feel lonely","scene":"support"}]',
                "application/json",
            )
        },
    )
    assert upload.status_code == 201

    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "Companion evaluation",
            "goal": "Evaluate conversation quality, tools and latency",
            "source_file_id": upload.json()["data"]["id"],
            "notification_targets": [{"channel": "email", "recipient": "reviewer@example.com"}],
        },
    )
    assert created.status_code == 201
    job_id = UUID(created.json()["data"]["id"])

    plan_agent_job(job_id)
    planned = client.get(f"/api/agent-jobs/{job_id}").json()["data"]
    assert planned["status"] == "completed"
    assert planned["eval_spec"]["source"]["fields"] == ["scene", "user_message"]
    assert planned["requires_approval"] is False

    steps = client.get(f"/api/agent-jobs/{job_id}/steps").json()["data"]
    assert [step["name"] for step in steps] == [
        "discover_source",
        "generate_eval_spec",
        "execute_eval_spec",
        "summarize",
    ]

    tasks = client.get("/api/human-tasks?task_status=pending").json()["data"]
    assert tasks == []
    assert planned["result"]["execution"]["execution_mode"] == "local_data"


def test_agent_assistant_returns_a_job_draft(client: TestClient, monkeypatch) -> None:
    login_as_developer(client)
    monkeypatch.setattr(
        "evalweave.api.routes.agents.assist_job_configuration",
        lambda *_: {
            "reply": "已识别接口和输出格式，可以创建任务。",
            "draft": {
                "title": "接口回复质量评测",
                "target_url": "https://model.test/chat",
                "target_body": {"message": "{{message}}"},
                "output_format": "xlsx",
            },
        },
    )
    response = client.post(
        "/api/agent/assist",
        json={"messages": [{"role": "user", "content": "评测这个接口并导出 Excel"}]},
    )
    assert response.status_code == 200
    assert response.json()["data"]["draft"]["output_format"] == "xlsx"


def test_source_parsing_options_accepts_model_field_mapping_shape() -> None:
    options = source_parsing_options(
        {
            "source": {
                "header_row": 2,
                "data_start_row": 3,
                "field_mappings": {
                    "student_name": "column_1（姓名）",
                    "student_id": "column_2（学号）",
                    "comprehensive_rank": "column_23（综测排名）",
                },
            }
        }
    )

    assert options[0:2] == (2, 3)
    assert options[2][0:2] == ["student_name", "student_id"]
    assert options[2][22] == "comprehensive_rank"


def test_result_output_formats() -> None:
    records = [
        {
            "case_index": 0,
            "status": "completed",
            "latency_ms": 12.5,
            "ttfb_ms": 4.5,
            "input": {"message": "你好"},
            "output": "您好",
            "evaluation": {
                "dimensions": [
                    {"key": "speed", "label": "速度评分", "score": 9, "reason": "响应较快"},
                    {"key": "relevance", "label": "匹配度评分", "score": 10, "reason": "回答匹配"},
                ],
                "overall_score": 9.5,
                "passed": True,
                "reason": "回答正确",
            },
        }
    ]
    excel, extension, _ = serialize_results(records, "xlsx")
    assert extension == "xlsx"
    workbook = load_workbook(BytesIO(excel), read_only=True)
    try:
        values = list(workbook.active.values)
        assert "综合评分" in values[0]
        assert "速度评分（1-10分）" in values[0]
        assert 9.5 in values[1]
        assert any("您好" in str(value) for value in values[1])
    finally:
        workbook.close()
    for output_format, expected_extension in [
        ("jsonl", "jsonl"),
        ("markdown", "md"),
        ("text", "txt"),
    ]:
        content, extension, _ = serialize_results(records, output_format)
        assert extension == expected_extension
        assert "您好" in content.decode()
        if output_format == "markdown":
            report = content.decode()
            assert "# AI 评测报告" in report
            assert "## 结论摘要" in report
            assert "## 核心指标" in report
            assert "## 代表性结果" in report
            assert "## 结果解读" in report
            assert "回答正确" in report


def test_tabular_markdown_is_a_readable_report_with_limited_preview() -> None:
    rows = [
        {
            "query": f"问题 {index}",
            "answer1": f"回答 A-{index}",
            "answer2": f"回答 B-{index}",
            "answer1_latency_ms": 100 + index,
            "answer2_latency_ms": 180 + index,
        }
        for index in range(8)
    ]
    summaries = [
        {
            "target": "接口 A",
            "total": 8,
            "succeeded": 8,
            "failed": 0,
            "average_latency_ms": 103.5,
            "p95_latency_ms": 107,
        },
        {
            "target": "接口 B",
            "total": 8,
            "succeeded": 7,
            "failed": 1,
            "average_latency_ms": 183.5,
            "p95_latency_ms": 187,
        },
    ]

    report = _tabular_markdown_report(rows, summaries, "双接口对比", "比较速度和回答")

    assert report.startswith("# 双接口对比")
    assert "## 结论摘要" in report
    assert "接口调用累计成功 **15/16** 次" in report
    assert "## 数据预览" in report
    assert "## 结果解读" in report
    assert "问题 4" in report
    assert "问题 5" not in report


def test_isolate_session_ids_replaces_nested_session_without_mutating_source() -> None:
    source = {"query": "1+1等于几", "session_id": "shared", "nested": [{"session_id": "shared"}]}
    isolated = isolate_session_ids(source)

    assert source["session_id"] == "shared"
    assert isolated["session_id"] != "shared"
    assert isolated["nested"][0]["session_id"] != "shared"
    assert isolated["session_id"] == isolated["nested"][0]["session_id"]
    assert isolate_session_ids(source)["session_id"] != isolated["session_id"]


def test_target_record_evaluation_accepts_dynamic_dimensions(monkeypatch) -> None:
    monkeypatch.setattr(
        "evalweave.agents.planner.request_json",
        lambda *_args, **_kwargs: {
            "evaluations": [
                {
                    "case_index": 0,
                    "dimensions": [
                        {"key": "speed", "label": "速度评分", "score": 8, "reason": "首包较快"}
                    ],
                    "overall_score": 8,
                    "passed": True,
                    "reason": "达到速度要求",
                }
            ]
        },
    )
    config = AgentConfig(enabled=True, base_url="https://model.test", model="judge")
    evaluations = evaluate_target_records(
        config,
        "对接口速度按1到10分评分",
        [
            {
                "case_index": 0,
                "status": "completed",
                "latency_ms": 850,
                "ttfb_ms": 120,
                "input": {"input": {"query": "你好"}, "expected": {"ttfb_ms_threshold": 500}},
                "output": "您好",
            }
        ],
    )

    assert evaluations[0]["dimensions"][0]["key"] == "speed"
    assert evaluations[0]["dimensions"][0]["score"] == 8


def test_agent_executes_http_target_without_host_allowlist(client: TestClient, monkeypatch) -> None:
    class FakeResponse:
        headers = {"content-type": "application/json"}

        def raise_for_status(self) -> None:
            return None

        def json(self):
            return {"data": {"reply": "I hear you"}}

    class FakeClient:
        def __init__(self, **_):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_):
            return None

        def post(self, *_args, **_kwargs):
            return FakeResponse()

    login_as_developer(client)
    project = current_project(client)
    upload = client.post(
        f"/api/projects/{project['id']}/files",
        files={"file": ("cases.json", b'[{"message":"hello"}]', "application/json")},
    ).json()["data"]
    monkeypatch.setattr("evalweave.agents.executor.httpx.Client", FakeClient)
    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "HTTP target",
            "goal": "Run target",
            "source_file_id": upload["id"],
            "requires_approval": False,
            "input_config": {
                "target": {
                    "url": "https://model.test/chat",
                    "body": {"message": "{{message}}"},
                    "response_path": "data.reply",
                }
            },
        },
    ).json()["data"]
    plan_agent_job(UUID(created["id"]))

    completed = client.get(f"/api/agent-jobs/{created['id']}").json()["data"]
    target = completed["result"]["execution"]["target"]
    assert target["executed"] is True
    assert target["succeeded"] == 1
    result_file = client.get(f"/api/files/{target['result_file_id']}/content")
    assert result_file.status_code == 200
    workbook = load_workbook(BytesIO(result_file.content), read_only=True)
    try:
        values = list(workbook.active.values)
        assert any("I hear you" in str(value) for value in values[1])
    finally:
        workbook.close()


def test_task_credentials_run_login_and_inject_token(client: TestClient) -> None:
    class LoginClient:
        cookies: dict[str, str] = {}

        def post(self, url: str, **kwargs):
            assert url == "https://auth.test/login"
            assert kwargs["json"] == {"username": "agent", "password": "secret"}
            return httpx.Response(
                200,
                json={"data": {"access_token": "runtime-token"}},
                request=httpx.Request("POST", url),
            )

    credentials = encrypt_secret_payload(
        {
            "headers": {"X-Task": "evaluation"},
            "auth": {
                "login_url": "https://auth.test/login",
                "body": {"username": "agent", "password": "secret"},
                "token_path": "data.access_token",
                "header_name": "Authorization",
                "header_prefix": "Bearer ",
            },
        }
    )

    _, headers = validate_target({"url": "https://target.test/chat", "credentials": credentials})
    authentication = authenticate_target(
        LoginClient(), "https://target.test/chat", headers, credentials
    )

    assert headers == {
        "X-Task": "evaluation",
        "Authorization": "Bearer runtime-token",
    }
    assert authentication["used"] is True
    assert authentication["source"] == "task"
    assert "runtime-token" not in str(authentication)


def test_python_workspace_registers_transformed_file(client: TestClient) -> None:
    login_as_developer(client)
    project = current_project(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["input", "output"])
    sheet.append(
        [
            json.dumps({"content": "你好"}, ensure_ascii=False),
            json.dumps({"answer": "您好"}, ensure_ascii=False),
        ]
    )
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    uploaded = client.post(
        f"/api/projects/{project['id']}/files",
        files={
            "file": (
                "source.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()["data"]
    code = """
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook

source = next(Path("inputs").glob("*.xlsx"))
workbook = load_workbook(source)
sheet = workbook.active
sheet["A2"] = json.loads(sheet["A2"].value)["content"]
sheet["B2"] = json.loads(sheet["B2"].value)["answer"]
workbook.save("outputs/cleaned.xlsx")
workbook.close()
"""

    with Session(get_engine()) as session:
        observation, updates = run_python_workspace(
            session,
            UUID(project["id"]),
            {"source_file_id": uploaded["id"]},
            code,
        )
        transformed = session.get(FileObject, UUID(updates["source_file_id"]))
        assert transformed is not None
        path = LocalFileStorage(get_settings().storage.local_directory).path_for(
            transformed.storage_key
        )
        rows = load_source_rows(path, transformed.original_name, 10)

    assert observation["outputs"][0]["file_name"] == "cleaned.xlsx"
    assert rows == [{"input": "你好", "output": "您好"}]


def test_python_workspace_creates_file_without_input(client: TestClient) -> None:
    login_as_developer(client)
    creator = client.get("/api/auth/me").json()["data"]
    project = current_project(client)
    code = """
from pathlib import Path

Path("outputs/generated.csv").write_text("query\\n你好\\n今天天气怎么样\\n", encoding="utf-8")
print("✅ 文件生成完成")
"""

    with Session(get_engine()) as session:
        observation, updates = run_python_workspace(
            session,
            UUID(project["id"]),
            {},
            code,
            source_file_ids=[],
            primary_output="generated.csv",
            created_by=UUID(creator["id"]),
        )
        generated = session.get(FileObject, UUID(updates["source_file_id"]))

    assert generated is not None
    assert generated.created_by == UUID(creator["id"])
    assert observation["outputs"][0]["file_name"] == "generated.csv"
    assert observation["primary_output_file_id"] == str(generated.id)
    assert observation["stdout"] == "✅ 文件生成完成"


def test_python_workspace_injects_selected_model_and_redacts_secret(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    login_as_developer(client)
    project = current_project(client)
    monkeypatch.setenv("UNRELATED_PRIVATE_VALUE", "must-not-be-inherited")
    code = """
import os

print(os.environ["EVALWEAVE_MODEL_BASE_URL"])
print(os.environ["EVALWEAVE_MODEL_NAME"])
print(os.environ["EVALWEAVE_MODEL_API_MODE"])
print(os.environ["EVALWEAVE_MODEL_API_KEY"])
print(os.environ.get("UNRELATED_PRIVATE_VALUE", "not-inherited"))
"""
    model_config = AgentConfig(
        enabled=True,
        api_mode="chat_completions",
        base_url="https://model.example/v1",
        model="followup-model",
        api_key="secret-model-key",
    )

    with Session(get_engine()) as session:
        observation, _ = run_python_workspace(
            session,
            UUID(project["id"]),
            {},
            code,
            source_file_ids=[],
            model_config=model_config,
        )

    assert observation["stdout"].splitlines() == [
        "https://model.example/v1",
        "followup-model",
        "chat_completions",
        "[REDACTED]",
        "not-inherited",
    ]


def test_python_workspace_rejects_output_containing_model_secret(
    client: TestClient,
) -> None:
    login_as_developer(client)
    project = current_project(client)
    model_config = AgentConfig(
        enabled=True,
        base_url="https://model.example/v1",
        model="followup-model",
        api_key="secret-model-key",
    )
    code = """
import os
from pathlib import Path

Path("outputs/leaked.txt").write_text(
    os.environ["EVALWEAVE_MODEL_API_KEY"],
    encoding="utf-8",
)
"""

    with (
        Session(get_engine()) as session,
        pytest.raises(
            ValueError,
            match="生成文件包含受保护的模型凭据",
        ),
    ):
        run_python_workspace(
            session,
            UUID(project["id"]),
            {},
            code,
            source_file_ids=[],
            model_config=model_config,
        )


def test_python_workspace_stops_running_process_when_cancelled(client: TestClient) -> None:
    login_as_developer(client)
    project = current_project(client)
    checks = 0

    def cancel_after_process_starts() -> None:
        nonlocal checks
        checks += 1
        raise RuntimeError("评测任务已取消")

    started = time.monotonic()
    with Session(get_engine()) as session, pytest.raises(RuntimeError, match="评测任务已取消"):
        run_python_workspace(
            session,
            UUID(project["id"]),
            {},
            "import time\ntime.sleep(30)\n",
            source_file_ids=[],
            timeout_seconds=None,
            cancellation_check=cancel_after_process_starts,
        )

    assert checks >= 1
    assert time.monotonic() - started < 5


def test_long_python_tool_creates_and_executes_background_job(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    login_as_developer(client)
    creator = client.get("/api/auth/me").json()["data"]
    project = current_project(client)
    queued: list[UUID] = []
    monkeypatch.setattr("evalweave.agents.tools.enqueue_python_job", queued.append)
    context = AssistantToolContext(
        draft={"title": "批量生成文件", "goal": "在后台生成结果文件"},
        project_id=UUID(project["id"]),
        actor_id=UUID(creator["id"]),
        trace=[
            {
                "name": "run_python",
                "label": "生成 Excel 前置文件",
                "status": "completed",
                "summary": {"stdout": "sensitive output is not copied"},
            },
            {
                "name": "submit_python_job",
                "label": "提交后台任务",
                "status": "running",
            },
        ],
    )

    result = json.loads(
        SubmitPythonJobTool().run(
            context,
            code=(
                "from pathlib import Path\n"
                'Path("outputs/result.txt").write_text("done", encoding="utf-8")\n'
            ),
            source_file_ids=[],
            primary_output="result.txt",
            evaluation_plan={
                "summary": "准备评测数据并生成整体结果",
                "steps": [
                    {
                        "title": "生成 Excel 前置文件",
                        "description": "在没有上传数据时生成评测所需的输入文件。",
                        "phase": "preparation",
                    },
                    {
                        "title": "执行回答质量评测",
                        "description": "逐条执行并记录评测结果。",
                        "phase": "execution",
                    },
                ],
            },
        )
    )

    job_id = UUID(result["job_id"])
    assert queued == [job_id]
    assert result["queued"] is True
    assert context.ui_action == {
        "type": "background_job",
        "job_id": str(job_id),
        "path": f"/evaluations/{job_id}",
    }
    with Session(get_engine()) as session:
        queued_steps = list(session.exec(select(AgentStep).where(AgentStep.job_id == job_id)))
    assert [step.status.value for step in queued_steps] == [
        "completed",
        "completed",
        "pending",
        "pending",
    ]

    observed_timeouts: list[float | None] = []
    summary_evidence: dict[str, object] = {}

    def capture_background_timeout(*args, **kwargs):
        observed_timeouts.append(kwargs.get("timeout_seconds"))
        return run_python_workspace(*args, **kwargs)

    monkeypatch.setattr(
        "evalweave.agents.workflow.run_python_workspace",
        capture_background_timeout,
    )

    def summarize_background_result(_config, _goal, execution, **_kwargs):
        summary_evidence.update(execution)
        return {"summary": "结果文件内容为 done，评测流程执行成功。"}

    monkeypatch.setattr(
        "evalweave.agents.workflow.generate_summary",
        summarize_background_result,
    )
    execute_python_job(job_id)

    with Session(get_engine()) as session:
        job = session.get(AgentJob, job_id)
        steps = list(session.exec(select(AgentStep).where(AgentStep.job_id == job_id)))
        output = session.get(FileObject, job.result_file_id) if job else None
    assert job is not None
    assert job.status.value == "completed"
    assert output is not None
    assert output.original_name == "result.txt"
    assert [step.name for step in steps] == [
        "generate_eval_spec",
        "prepare_data",
        "execute_eval_spec",
        "summarize",
    ]
    assert steps[1].output_data == {
        "label": "准备评测数据",
        "phase": "preparation",
        "source_file_count": 0,
    }
    assert job.eval_spec["data_program"]["steps"][0]["title"] == "生成 Excel 前置文件"
    assert job.result["execution"]["execution_mode"] == "agent_python"
    assert job.result["summary"] == "结果文件内容为 done，评测流程执行成功。"
    assert summary_evidence["primary_result"]["content"] == "done"
    assert observed_timeouts == [None]


def test_dialog_python_tool_uses_30_second_timeout(monkeypatch: pytest.MonkeyPatch) -> None:
    observed_timeouts: list[float | None] = []

    def capture_dialog_timeout(*_args, **kwargs):
        observed_timeouts.append(kwargs.get("timeout_seconds"))
        return ({"stdout": "", "outputs": []}, {})

    monkeypatch.setattr(
        "evalweave.agents.tools.run_python_workspace",
        capture_dialog_timeout,
    )
    context = AssistantToolContext(
        draft={},
        project_id=UUID(int=1),
        actor_id=UUID(int=2),
    )

    result = json.loads(
        RunPythonTool().run(
            context,
            code="print('done')",
            source_file_ids=[],
        )
    )

    assert result["ok"] is True
    assert observed_timeouts == [30]


def test_agent_can_inspect_existing_background_job_before_recreating(
    client: TestClient,
) -> None:
    login_as_developer(client)
    creator = client.get("/api/auth/me").json()["data"]
    project = current_project(client)
    with Session(get_engine()) as session:
        job = AgentJob(
            project_id=UUID(project["id"]),
            created_by=UUID(creator["id"]),
            title="十轮连续对话评测",
            goal="执行十轮连续对话并生成 Excel",
            status=AgentJobStatus.FAILED,
            input_config={"job_type": "python", "python_code": "raise TimeoutError()"},
            error="目标接口连续三次读取超时",
            repair_attempts=3,
            max_repair_attempts=3,
        )
        session.add(job)
        session.commit()
        session.refresh(job)
        session.add(
            AgentStep(
                job_id=job.id,
                name="execute_eval_spec",
                status=StepStatus.FAILED,
                attempt=3,
                error="ReadTimeout",
            )
        )
        session.commit()
        job_id = job.id

    context = AssistantToolContext(
        draft={"background_job_id": str(job_id)},
        project_id=UUID(project["id"]),
        actor_id=UUID(creator["id"]),
    )
    result = json.loads(InspectAgentJobTool().run(context))

    assert result["job_id"] == str(job_id)
    assert result["status"] == "failed"
    assert result["error"] == "目标接口连续三次读取超时"
    assert result["steps"][0]["status"] == "failed"
    assert result["repair_attempts"] == 3


def test_background_python_job_repairs_script_before_failing(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    login_as_developer(client)
    creator = client.get("/api/auth/me").json()["data"]
    project = current_project(client)
    with Session(get_engine()) as session:
        job = AgentJob(
            project_id=UUID(project["id"]),
            created_by=UUID(creator["id"]),
            title="批量接口评测",
            goal="超时数据记录失败后继续执行",
            status=AgentJobStatus.PENDING,
            input_config={
                "job_type": "python",
                "python_code": "raise TimeoutError('timed out')",
                "source_file_ids": [],
                "output_format": "file",
            },
            eval_spec={},
            requires_approval=False,
            max_repair_attempts=2,
        )
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    executed_codes: list[str] = []

    def fake_workspace(_session, _project_id, _draft, code, *_args, **_kwargs):
        executed_codes.append(code)
        if len(executed_codes) == 1:
            raise ValueError("ReadTimeout: timed out")
        return ({"stdout": "已跳过 1 条超时数据", "outputs": []}, {})

    monkeypatch.setattr("evalweave.agents.workflow.run_python_workspace", fake_workspace)
    monkeypatch.setattr(
        "evalweave.agents.workflow.resolve_agent_config",
        lambda *_args: AgentConfig(
            enabled=True,
            base_url="https://model.test/v1",
            model="repair-model",
        ),
    )
    monkeypatch.setattr(
        "evalweave.agents.workflow.repair_python_script",
        lambda _config, _goal, _code, _error: "print('continue after timeout')",
    )
    monkeypatch.setattr(
        "evalweave.agents.workflow.generate_summary",
        lambda *_args, **_kwargs: {"summary": "执行完成，1 条超时数据已记录为失败。"},
    )

    execute_python_job(job_id)

    with Session(get_engine()) as session:
        repaired_job = session.get(AgentJob, job_id)
        steps = list(
            session.exec(
                select(AgentStep).where(AgentStep.job_id == job_id).order_by(AgentStep.created_at)
            )
        )

    assert repaired_job is not None
    assert repaired_job.status == AgentJobStatus.COMPLETED
    assert repaired_job.repair_attempts == 1
    assert repaired_job.input_config["python_code"] == "print('continue after timeout')"
    assert executed_codes == [
        "raise TimeoutError('timed out')",
        "print('continue after timeout')",
    ]
    assert [step.status for step in steps if step.name == "execute_eval_spec"] == [
        StepStatus.COMPLETED,
    ]
    assert next(step for step in steps if step.name == "execute_eval_spec").attempt == 2


def test_worker_restart_recovers_interrupted_python_job(client: TestClient) -> None:
    login_as_developer(client)
    creator = client.get("/api/auth/me").json()["data"]
    project = current_project(client)
    with Session(get_engine()) as session:
        job = AgentJob(
            project_id=UUID(project["id"]),
            created_by=UUID(creator["id"]),
            title="中断任务",
            goal="恢复中断任务",
            status=AgentJobStatus.RUNNING,
            input_config={"job_type": "python", "python_code": "print('ok')"},
        )
        session.add(job)
        session.commit()
        session.refresh(job)
        step = AgentStep(
            job_id=job.id,
            name="execute_eval_spec",
            status=StepStatus.RUNNING,
        )
        session.add(step)
        session.commit()

        recovered = recover_interrupted_python_jobs(session)
        session.refresh(job)
        session.refresh(step)

        assert recovered == [job.id]
        assert job.status == AgentJobStatus.PENDING
        assert step.status == StepStatus.PENDING
        assert step.started_at is None


def test_python_job_repair_reuses_original_conversation_context(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    login_as_developer(client)
    creator = client.get("/api/auth/me").json()["data"]
    project = current_project(client)
    captured: dict[str, object] = {}
    with Session(get_engine()) as session:
        conversation = AssistantConversation(
            project_id=UUID(project["id"]),
            created_by=UUID(creator["id"]),
            draft={"output_format": "xlsx", "concurrency": 5},
        )
        session.add(conversation)
        session.commit()
        session.refresh(conversation)
        session.add(
            AssistantMessage(
                conversation_id=conversation.id,
                role="user",
                content="生成 50 条日常问答，并发 5，结果保存为中文名 Excel",
            )
        )
        job = AgentJob(
            project_id=UUID(project["id"]),
            created_by=UUID(creator["id"]),
            title="日常问答评测",
            goal="生成并评测 50 条日常问答",
            input_config={
                "job_type": "python",
                "conversation_id": str(conversation.id),
                "python_code": "requests.post(url, timeout=60)",
                "output_format": "xlsx",
                "primary_output": "日常问答评测.xlsx",
                "source_file_ids": [],
            },
            eval_spec={"summary": "批量调用接口"},
        )
        session.add(job)
        session.commit()
        session.refresh(job)
        conversation_id = conversation.id
        job_id = job.id

        def fake_react(*args, **_kwargs):
            captured["messages"] = args[1]
            captured["workspace"] = args[2]
            captured["conversation_id"] = args[5]
            captured["repair_job_id"] = args[6]
            captured["allowed_tools"] = args[7]
            job.input_config = {
                **job.input_config,
                "python_code": "# repaired\ncontinue_on_timeout()",
            }
            session.add(job)
            session.commit()
            yield (
                "result",
                {"ui_action": {"type": "python_repaired", "summary": "已修复超时处理"}},
            )

        monkeypatch.setattr(
            "evalweave.agents.workflow.stream_react_configuration",
            fake_react,
        )
        repaired = repair_python_job_with_react(
            session,
            job,
            AgentConfig(enabled=True, base_url="https://model.test/v1", model="repair"),
            "requests.post(url, timeout=60)",
            TimeoutError("Read timed out"),
            1,
        )

    assert repaired == "# repaired\ncontinue_on_timeout()"
    assert captured["conversation_id"] == conversation_id
    assert captured["repair_job_id"] == job_id
    assert captured["allowed_tools"] == {
        "repair_python_job_script",
        "inspect_source",
        "probe_http_target",
        "run_python",
    }
    messages = captured["messages"]
    assert isinstance(messages, list)
    assert "生成 50 条日常问答" in messages[0]["content"]
    assert "Read timed out" in messages[-1]["content"]
    workspace = captured["workspace"]
    assert isinstance(workspace, dict)
    assert workspace["current_draft"]["concurrency"] == 5
    assert workspace["execution_repair"]["python_code"] == "requests.post(url, timeout=60)"
    assert workspace["execution_repair"]["primary_output"] == "日常问答评测.xlsx"


def test_agent_generates_cases_when_http_target_has_no_source_file(
    client: TestClient, monkeypatch
) -> None:
    class FakeResponse:
        headers = {"content-type": "application/json"}

        def raise_for_status(self) -> None:
            return None

        def json(self):
            return {"reply": "ok"}

    class FakeClient:
        def __init__(self, **_):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_):
            return None

        def post(self, *_args, **_kwargs):
            return FakeResponse()

    login_as_developer(client)
    project = current_project(client)
    monkeypatch.setattr("evalweave.agents.executor.httpx.Client", FakeClient)
    monkeypatch.setattr(
        "evalweave.agents.workflow.generate_test_cases",
        lambda *_, **__: [
            {"input": {"message": "hello"}, "expected": {}, "metadata": {"category": "normal"}},
            {
                "input": {"message": "edge case"},
                "expected": {},
                "metadata": {"category": "edge"},
            },
        ],
    )
    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "Generated HTTP cases",
            "goal": "Generate two chat cases and run the target",
            "requires_approval": False,
            "input_config": {
                "max_cases": 2,
                "target": {
                    "url": "https://model.test/chat",
                    "body": {"message": "{{message}}"},
                },
            },
        },
    ).json()["data"]
    plan_agent_job(UUID(created["id"]))

    completed = client.get(f"/api/agent-jobs/{created['id']}").json()["data"]
    target = completed["result"]["execution"]["target"]
    assert completed["status"] == "completed"
    assert target["case_source"] == "ai_generated"
    assert target["total_cases"] == 2
    assert completed["eval_spec"]["source"]["case_count"] == 2
    steps = client.get(f"/api/agent-jobs/{created['id']}/steps").json()["data"]
    assert [step["name"] for step in steps] == [
        "generate_eval_spec",
        "generate_test_cases",
        "validate_target",
        "execute_eval_spec",
        "summarize",
    ]
    event_stream = client.get(f"/api/agent-jobs/{created['id']}/events")
    assert event_stream.status_code == 200
    assert "event: agent-event" in event_stream.text
    assert "event: job-state" in event_stream.text
    assert '"steps":' in event_stream.text
    assert '"event_type": "model_start"' in event_stream.text
    assert "event: end" in event_stream.text


def test_target_preflight_failure_returns_to_react_and_retries(
    client: TestClient, monkeypatch
) -> None:
    login_as_developer(client)
    project = current_project(client)
    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "Repair target request",
            "goal": "Validate a chat endpoint",
            "requires_approval": False,
            "input_config": {
                "target": {
                    "url": "https://model.test/chat",
                    "body": '{"query":"{{generated_query}}"}',
                }
            },
        },
    ).json()["data"]
    attempts = 0

    def fake_validate(_session, _job):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise ValueError("HTTP 422; response body: expected an object")
        return {"validated": True, "records": [], "response_modes": ["streaming"]}

    def fake_react(*_args):
        running = {"name": "update_task_draft", "status": "running"}
        yield "tool_start", running
        yield (
            "tool_result",
            {
                **running,
                "status": "completed",
                "summary": {"ok": True, "updated_fields": ["target_body"]},
            },
        )
        yield (
            "result",
            {
                "draft": {
                    "target_url": "https://model.test/chat",
                    "target_body": {"query": "{{generated_query}}"},
                    "target_validated": True,
                    "expected_streaming": True,
                },
                "ui_action": None,
            },
        )

    monkeypatch.setattr("evalweave.agents.workflow.validate_http_target", fake_validate)
    monkeypatch.setattr("evalweave.agents.workflow.stream_react_configuration", fake_react)
    config = AgentConfig(enabled=True, base_url="https://model.test", model="test")
    with Session(get_engine()) as session:
        job = session.get(AgentJob, UUID(created["id"]))
        assert job is not None
        job.eval_spec = {"generated_cases": [{"input": {"generated_query": "hello"}}]}
        session.add(job)
        session.commit()
        result = validate_http_target_with_react(session, job, config)
        session.refresh(job)
        assert result["validated"] is True
        assert job.repair_attempts == 1
        assert job.input_config["target"]["body"] == {"query": "{{generated_query}}"}
    steps = client.get(f"/api/agent-jobs/{created['id']}/steps").json()["data"]
    assert [(step["name"], step["status"]) for step in steps] == [
        ("validate_target", "failed"),
        ("update_task_draft", "completed"),
        ("validate_target", "completed"),
    ]


def test_semantic_preflight_rejects_when_all_sample_responses_are_wrong(
    client: TestClient, monkeypatch
) -> None:
    login_as_developer(client)
    project = current_project(client)
    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "Validate response meaning",
            "goal": "检查回答是否正确",
            "requires_approval": False,
            "input_config": {"target": {"url": "https://model.test/chat"}},
        },
    ).json()["data"]
    records = [
        {
            "case_index": 0,
            "status": "completed",
            "input": {"input": {"content": "太阳系行星顺序"}, "expected": {"correct": True}},
            "output": {"code": 200, "data": "错误答案"},
        },
        {
            "case_index": 1,
            "status": "completed",
            "input": {"input": {"content": "计算价格"}, "expected": {"correct": True}},
            "output": {"code": 200, "data": "无法回答"},
        },
    ]
    monkeypatch.setattr(
        "evalweave.agents.workflow.validate_http_target",
        lambda *_: {"validated": True, "records": records},
    )
    monkeypatch.setattr(
        "evalweave.agents.workflow.evaluate_target_records",
        lambda *_args, **_kwargs: [
            {
                "case_index": record["case_index"],
                "passed": False,
                "reason": "回答与问题和预期不符",
            }
            for record in records
        ],
    )
    monkeypatch.setattr(
        "evalweave.agents.workflow.run_streamed_model_output",
        lambda _job, _phase, _label, operation: operation(lambda _delta: None),
    )
    config = AgentConfig(enabled=True, base_url="https://judge.test", model="judge")

    with Session(get_engine()) as session:
        job = session.get(AgentJob, UUID(created["id"]))
        assert job is not None
        with pytest.raises(ValueError, match="抽检响应全部不符合测试预期"):
            validate_http_target_semantics(session, job, config)


def test_parse_target_response_detects_server_sent_events() -> None:
    class StreamingResponse:
        headers = {"content-type": "text/event-stream; charset=utf-8"}
        text = 'data: {"delta":"hello"}\n\ndata: {"delta":" world"}\n\ndata: [DONE]\n'

    output, response_mode = parse_target_response(StreamingResponse(), "")

    assert response_mode == "streaming"
    assert output == {"events": [{"delta": "hello"}, {"delta": " world"}]}


def test_parse_target_response_supports_stream_wildcard_path() -> None:
    class StreamingResponse:
        headers = {"content-type": "text/event-stream"}
        text = (
            'data: {"data":{"message":"hello"}}\n\n'
            'data: {"data":{"message":" world"}}\n\n'
            "data: [DONE]\n"
        )

    output, response_mode = parse_target_response(StreamingResponse(), "events[*].data.message")

    assert response_mode == "streaming"
    assert output == "hello world"


def test_normalize_target_body_decodes_json_object_template() -> None:
    body = normalize_target_body('{"query":"{{generated_query}}","plugins":[]}')

    assert body == {"query": "{{generated_query}}", "plugins": []}


def test_session_id_isolation_supports_camel_and_snake_case() -> None:
    isolated = isolate_session_ids(
        {
            "sessionId": "shared-camel",
            "nested": {"session_id": "shared-snake"},
        },
        session_id="isolated-case",
    )

    assert isolated["sessionId"] == "isolated-case"
    assert isolated["nested"]["session_id"] == "isolated-case"


def test_target_case_rejects_application_error_inside_http_200() -> None:
    class FakeClient:
        def post(self, *_args, **_kwargs) -> httpx.Response:
            return httpx.Response(
                200,
                headers={"content-type": "application/json"},
                json={"code": 500, "msg": "服务器内部错误!", "data": None},
                request=httpx.Request("POST", "https://model.test/chat"),
            )

    record = execute_target_case(
        FakeClient(),
        "https://model.test/chat",
        {},
        {"body": "{{row}}"},
        {"input": {"sessionId": "shared", "content": "hello"}},
        0,
    )

    assert record["status"] == "failed"
    assert record["status_code"] == 200
    assert "code=500" in record["error"]
    assert "服务器内部错误" in record["response_body"]


def test_application_error_detector_allows_success_envelope() -> None:
    response = httpx.Response(
        200,
        headers={"content-type": "application/json"},
        json={"code": 200, "msg": "ok", "data": {"answer": "hello"}},
    )

    assert detect_application_error(response) is None


def test_assistant_conversation_can_be_renamed_and_deleted(client: TestClient) -> None:
    login_as_developer(client)
    project = current_project(client)
    created = client.post("/api/assistant/conversations", json={"project_id": project["id"]})
    conversation_id = created.json()["data"]["id"]

    renamed = client.patch(
        f"/api/assistant/conversations/{conversation_id}",
        json={"title": "  新的对话名称  "},
    )

    assert renamed.status_code == 200
    assert renamed.json()["data"]["title"] == "新的对话名称"
    assert client.delete(f"/api/assistant/conversations/{conversation_id}").status_code == 204
    assert client.get(f"/api/assistant/conversations/{conversation_id}/messages").status_code == 404
    conversations = client.get("/api/assistant/conversations").json()["data"]
    assert all(item["id"] != conversation_id for item in conversations)


def test_conversation_title_is_normalized_to_two_to_ten_characters() -> None:
    assert (
        normalize_conversation_title("标题：批量接口质量评测方案", "帮我测试接口")
        == "批量接口质量评测方案"
    )
    assert normalize_conversation_title("？", "天气") == "天气"
    assert normalize_conversation_title("", "啊") == "新对话"


def test_assistant_conversation_streams_and_persists_draft(client: TestClient, monkeypatch) -> None:
    login_as_developer(client)
    project = current_project(client)
    conversation = client.post("/api/assistant/conversations", json={"project_id": project["id"]})
    assert conversation.status_code == 201
    conversation_id = conversation.json()["data"]["id"]
    greeting = client.get(f"/api/assistant/conversations/{conversation_id}/messages").json()["data"]
    assert greeting[0]["role"] == "assistant"

    uploaded = client.post(
        f"/api/projects/{project['id']}/files",
        data={"category": "dataset_source"},
        files={
            "file": (
                "cases.xlsx",
                b"test workbook",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201
    scheduled_titles: list[tuple[object, ...]] = []
    monkeypatch.setattr(
        "evalweave.api.routes.agents.schedule_conversation_title",
        lambda *args: scheduled_titles.append(args),
    )

    monkeypatch.setattr(
        "evalweave.api.routes.agents.assist_job_configuration",
        lambda *_: {
            "reply": "已收到需求。",
            "draft": {
                "title": "成绩比拼",
                "goal": "比较综合成绩和各项排名",
                "target_url": "https://model.test/scores",
                "target_body": {"score": "{{score}}"},
                "target_validated": True,
                "auth_required": True,
            },
        },
    )
    response = client.post(
        f"/api/assistant/conversations/{conversation_id}/messages/stream",
        json={
            "content": "比较成绩",
            "source_file_id": uploaded.json()["data"]["id"],
        },
    )
    events = [json.loads(line) for line in response.text.splitlines()]
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-cache, no-transform"
    assert response.headers["x-accel-buffering"] == "no"
    assert events[0]["type"] == "start"
    assert events[-1]["type"] == "done"
    assert events[-1]["stage"] == "choose_output"
    assert events[-1]["content"] == "已收到需求。"
    assert len(scheduled_titles) == 1
    assert scheduled_titles[0][2] == "比较成绩"

    stored = client.get("/api/assistant/conversations").json()["data"][0]
    assert stored["draft"]["goal"] == "比较综合成绩和各项排名"
    assert stored["status"] == "choose_output"
    messages = client.get(f"/api/assistant/conversations/{conversation_id}/messages").json()["data"]
    assert [message["role"] for message in messages] == [
        "assistant",
        "user",
        "assistant",
    ]
    assert messages[1]["attachment_file_id"] == uploaded.json()["data"]["id"]
    assert messages[1]["attachment_name"] == "cases.xlsx"
    assert messages[1]["attachment_content_type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert messages[1]["attachment_size_bytes"] == len(b"test workbook")

    monkeypatch.setattr(
        "evalweave.api.routes.agents.assist_job_configuration",
        lambda *_: {
            "reply": "将使用上传的数据比较成绩并生成 Excel 结果，是否开始执行？",
            "draft": {},
            "ui_action": {"type": "confirm"},
        },
    )
    ready_response = client.post(
        f"/api/assistant/conversations/{conversation_id}/messages/stream",
        json={"content": "Excel 文件", "output_format": "xlsx"},
    )
    ready_events = [json.loads(line) for line in ready_response.text.splitlines()]
    assert ready_events[-1]["stage"] == "ready"
    assert ready_events[-1]["draft"]["output_format"] == "xlsx"
    assert ready_events[-1]["content"] == (
        "将使用上传的数据比较成绩并生成 Excel 结果，是否开始执行？"
    )
    assert len(scheduled_titles) == 1

    job = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "成绩比拼",
            "goal": "比较综合成绩和各项排名",
            "output_format": "xlsx",
            "input_config": {"target": {"url": "https://model.test/scores"}},
        },
    ).json()["data"]
    started = client.post(
        f"/api/assistant/conversations/{conversation_id}/started",
        json={"agent_job_id": job["id"]},
    )
    assert started.status_code == 200
    assert started.json()["data"]["status"] == "started"
    assert started.json()["data"]["agent_job_id"] == job["id"]
    messages = client.get(f"/api/assistant/conversations/{conversation_id}/messages").json()["data"]
    assert messages[-1]["role"] == "assistant"
    assert messages[-1]["content"] == ("将使用上传的数据比较成绩并生成 Excel 结果，是否开始执行？")


def test_assistant_stream_forwards_react_tool_events(client: TestClient, monkeypatch) -> None:
    login_as_developer(client)
    project = current_project(client)
    conversation = client.post(
        "/api/assistant/conversations", json={"project_id": project["id"]}
    ).json()["data"]
    config = AgentConfig(enabled=True, base_url="https://model.test", model="test-model")
    monkeypatch.setattr("evalweave.api.routes.agents.resolve_agent_config", lambda *_: config)

    def fake_react(*_):
        running = {"name": "probe_http_target", "label": "验证目标接口", "status": "running"}
        completed = {**running, "status": "completed", "summary": {"status_code": 200}}
        yield "delta", "The endpoint is reachable. "
        yield "round_end", None
        yield "tool_start", running
        yield "tool_result", completed
        yield "delta", "Choose a delivery format."
        yield (
            "result",
            {
                "reply": "Choose a delivery format.",
                "draft": {
                    "title": "接口评测",
                    "goal": "验证接口质量",
                    "target_url": "https://model.test/chat",
                    "target_body": {"query": "{{query}}"},
                    "target_validated": True,
                },
                "ui_action": {"type": "choose_output"},
                "react_trace": [completed],
            },
        )

    monkeypatch.setattr("evalweave.api.routes.agents.stream_react_configuration", fake_react)
    response = client.post(
        f"/api/assistant/conversations/{conversation['id']}/messages/stream",
        json={"content": "评测这个接口"},
    )
    events = [json.loads(line) for line in response.text.splitlines()]

    assert [event["type"] for event in events] == [
        "start",
        "delta",
        "round_end",
        "tool_start",
        "tool_result",
        "delta",
        "done",
    ]
    assert events[-1]["content"] == "Choose a delivery format."
    assert events[-1]["stage"] == "choose_output"
    assert events[-1]["ui_action"] == {"type": "choose_output"}
    assert events[-1]["draft"]["react_trace"][0]["status"] == "completed"
    assert "ui_action" not in events[-1]["draft"]
    stored_messages = client.get(
        f"/api/assistant/conversations/{conversation['id']}/messages"
    ).json()["data"]
    assert stored_messages[-2]["content"] == "The endpoint is reachable."
    assert stored_messages[-2]["include_in_context"] is False
    assert stored_messages[-1]["content"] == events[-1]["content"]
    assert stored_messages[-1]["ui_action"] == {"type": "choose_output"}
    assert stored_messages[-1]["include_in_context"] is True
    assert stored_messages[-1]["is_streaming"] is False

    snapshot_response = client.get(
        f"/api/assistant/conversations/{conversation['id']}/messages/events"
    )
    snapshot_line = next(
        line for line in snapshot_response.text.splitlines() if line.startswith("data: ")
    )
    snapshot = json.loads(snapshot_line.removeprefix("data: "))
    assert snapshot[-2]["include_in_context"] is False
    assert snapshot[-1]["is_streaming"] is False

    captured_history = []

    def capture_react(_config, messages, *_args):
        captured_history.extend(messages)
        yield (
            "result",
            {
                "reply": "Next final reply.",
                "draft": {},
                "ui_action": None,
                "react_trace": [],
            },
        )

    monkeypatch.setattr("evalweave.api.routes.agents.stream_react_configuration", capture_react)
    client.post(
        f"/api/assistant/conversations/{conversation['id']}/messages/stream",
        json={"content": "Continue"},
    )
    context_contents = [message["content"] for message in captured_history]
    assert "The endpoint is reachable." not in context_contents
    assert "Choose a delivery format." in context_contents


def test_assistant_user_input_keeps_conversation_collecting(
    client: TestClient, monkeypatch
) -> None:
    login_as_developer(client)
    project = current_project(client)
    conversation = client.post(
        "/api/assistant/conversations", json={"project_id": project["id"]}
    ).json()["data"]
    config = AgentConfig(enabled=True, base_url="https://model.test", model="test-model")
    monkeypatch.setattr("evalweave.api.routes.agents.resolve_agent_config", lambda *_: config)

    def fake_react(*_):
        yield (
            "result",
            {
                "reply": "Please provide the missing request field.",
                "draft": {
                    "title": "API evaluation",
                    "goal": "Evaluate answer quality",
                    "target_url": "https://model.test/chat",
                    "target_body": {"query": "{{query}}"},
                    "target_validated": True,
                    "output_format": "xlsx",
                },
                "ui_action": {
                    "type": "user_input",
                    "question": "Please provide the missing request field.",
                    "options": [],
                },
                "react_trace": [],
            },
        )

    monkeypatch.setattr("evalweave.api.routes.agents.stream_react_configuration", fake_react)
    response = client.post(
        f"/api/assistant/conversations/{conversation['id']}/messages/stream",
        json={"content": "Continue configuring"},
    )
    events = [json.loads(line) for line in response.text.splitlines()]

    assert events[-1]["stage"] == "collecting"


def test_assistant_completed_tool_reply_does_not_request_task_confirmation(
    client: TestClient, monkeypatch
) -> None:
    login_as_developer(client)
    project = current_project(client)
    conversation = client.post(
        "/api/assistant/conversations", json={"project_id": project["id"]}
    ).json()["data"]
    generated_file = client.post(
        f"/api/projects/{project['id']}/files",
        data={"category": "dataset_source"},
        files={"file": ("scored_results.xlsx", b"generated workbook", "application/xlsx")},
    ).json()["data"]
    config = AgentConfig(enabled=True, base_url="https://model.test", model="test-model")
    monkeypatch.setattr("evalweave.api.routes.agents.resolve_agent_config", lambda *_: config)

    def fake_react(*_):
        yield (
            "result",
            {
                "reply": "文件已经重新评分并生成，可以直接下载。",
                "draft": {
                    "title": "模型回答质量评分",
                    "goal": "重新评估回答相关性",
                    "task_mode": "local_analysis",
                    "source_file_id": generated_file["id"],
                    "source_inspected": True,
                    "output_format": "xlsx",
                },
                "ui_action": None,
                "react_trace": [
                    {
                        "name": "run_python",
                        "label": "运行 Python 文件处理",
                        "status": "completed",
                        "summary": {
                            "ok": True,
                            "primary_output_file_id": generated_file["id"],
                            "outputs": [
                                {
                                    "file_id": generated_file["id"],
                                    "file_name": "scored_results.xlsx",
                                }
                            ],
                        },
                    }
                ],
            },
        )

    monkeypatch.setattr("evalweave.api.routes.agents.stream_react_configuration", fake_react)
    response = client.post(
        f"/api/assistant/conversations/{conversation['id']}/messages/stream",
        json={"content": "重新评分这个文件"},
    )
    events = [json.loads(line) for line in response.text.splitlines()]

    assert events[-1]["type"] == "done", events
    assert events[-1]["stage"] == "collecting"
    assert events[-1]["ui_action"] is None
    assert "ui_action" not in events[-1]["draft"]
    assert events[-1]["attachment_file_id"] == generated_file["id"]
    response = client.get(f"/api/assistant/conversations/{conversation['id']}/messages")
    messages = response.json()["data"]
    assert messages[-1]["ui_action"] is None
    assert messages[-1]["attachment_file_id"] == generated_file["id"]
    assert messages[-1]["attachment_name"] == "scored_results.xlsx"


def test_generic_data_program_combines_model_and_multiple_http_targets(
    client: TestClient, monkeypatch
) -> None:
    login_as_developer(client)
    project = current_project(client)
    source = client.post(
        f"/api/projects/{project['id']}/files",
        files={
            "file": (
                "queries.json",
                b'[{"query":"hello"},{"query":"edge case"}]',
                "application/json",
            )
        },
    ).json()["data"]
    created = client.post(
        f"/api/projects/{project['id']}/agent-jobs",
        json={
            "title": "Generic enrichment",
            "goal": "Generate cases and compare two APIs",
            "source_file_id": source["id"],
            "output_format": "xlsx",
            "requires_approval": False,
            "input_config": {
                "targets": [
                    {
                        "name": "API 1",
                        "url": "https://model.test/one",
                        "body": {"query": "{{query}}"},
                    },
                    {
                        "name": "API 2",
                        "url": "https://model.test/two",
                        "body": {"query": "{{query}}"},
                    },
                ]
            },
        },
    ).json()["data"]
    monkeypatch.setattr(
        "evalweave.agents.executor.authenticate_target", lambda *_args, **_kwargs: "none"
    )

    def fake_target_case(_client, url, _headers, _target, row, index):
        target_number = 1 if url.endswith("/one") else 2
        return {
            "case_index": index,
            "status": "completed",
            "latency_ms": float(100 * target_number + index),
            "ttfb_ms": float(50 * target_number + index),
            "input": row,
            "output": f"answer-{target_number}-{index}",
            "response_mode": "json",
        }

    monkeypatch.setattr("evalweave.agents.executor.execute_target_case", fake_target_case)
    with Session(get_engine()) as session:
        job = session.get(AgentJob, UUID(created["id"]))
        assert job is not None
        job.eval_spec = {
            "source": {
                "header_row": 1,
                "field_mappings": {"query": "canonical_query"},
            },
            "data_program": {
                "steps": [
                    {"primitive": "convert"},
                    {
                        "primitive": "model_map",
                        "instruction": "Generate two additional cases",
                        "input_fields": ["canonical_query"],
                        "output_columns": [{"name": "generated_cases", "type": "array"}],
                    },
                    {"action": "http_map", "target_index": 0, "answer_column": "answer1"},
                    {"action": "http_map", "target_index": 1, "answer_column": "answer2"},
                    {"primitive": "aggregate", "instruction": "Summarize numeric metrics"},
                    {"primitive": "summarize"},
                    {"primitive": "format_convert"},
                ]
            },
        }
        session.add(job)
        session.commit()

        result = execute_data_program(
            session,
            job,
            lambda _instruction, rows, _columns, _index: [
                {
                    "generated_cases": [
                        f"{row['canonical_query']}-normal",
                        f"{row['canonical_query']}-edge",
                    ]
                }
                for row in rows
            ],
        )
        result_file_id = result["result_file_id"]

    assert result["summaries"][0]["average_latency_ms"] == 100.5
    assert result["summaries"][1]["average_latency_ms"] == 200.5
    download = client.get(f"/api/files/{result_file_id}/content")
    workbook = load_workbook(BytesIO(download.content), read_only=True)
    try:
        rows = list(workbook["结果"].values)
        assert "generated_cases" in rows[0]
        assert "answer1" in rows[0]
        assert "answer2" in rows[0]
        assert "hello-normal" in rows[1][rows[0].index("generated_cases")]
        assert workbook["汇总"].max_row == 4
    finally:
        workbook.close()
