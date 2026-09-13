from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, list | tuple):
        return [json_safe(item) for item in value]
    return str(value)


def inspect_source(path: Path, original_name: str, sample_size: int) -> dict[str, Any]:
    extension = Path(original_name).suffix.lower()
    if extension == ".json":
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = value if isinstance(value, list) else [value]
        return summarize_rows("json", rows, sample_size)
    if extension == ".jsonl":
        rows = []
        with path.open("r", encoding="utf-8-sig") as stream:
            for line in stream:
                if line.strip():
                    rows.append(json.loads(line))
                if len(rows) >= sample_size:
                    break
        return summarize_rows("jsonl", rows, sample_size, sampled_only=True)
    if extension == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) >= sample_size:
                    break
            return summarize_rows(
                "csv", rows, sample_size, fields=reader.fieldnames or [], sampled_only=True
            )
    if extension == ".xlsx":
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                headers = [
                    str(value or f"column_{index + 1}")
                    for index, value in enumerate(next(iterator, []))
                ]
                rows = []
                for values in iterator:
                    rows.append(dict(zip(headers, values, strict=False)))
                    if len(rows) >= sample_size:
                        break
                result = summarize_rows(
                    "xlsx", rows, sample_size, fields=headers, sampled_only=True
                )
                result["sheet"] = sheet.title
                return result
            finally:
                workbook.close()
    raise ValueError(f"Unsupported source format: {extension or '(none)'}")


def load_source_rows(
    path: Path,
    original_name: str,
    max_cases: int,
    *,
    header_row: int = 1,
    data_start_row: int | None = None,
    field_names: list[str] | None = None,
) -> list[dict[str, Any]]:
    extension = Path(original_name).suffix.lower()
    rows: list[Any]
    if extension == ".json":
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = value if isinstance(value, list) else [value]
    elif extension == ".jsonl":
        rows = []
        with path.open("r", encoding="utf-8-sig") as stream:
            for line in stream:
                if line.strip():
                    rows.append(json.loads(line))
                if len(rows) >= max_cases:
                    break
    elif extension == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))[:max_cases]
    elif extension == ".xlsx":
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                raw_rows = list(
                    sheet.iter_rows(max_row=max_cases * 10 + 100, values_only=True)
                )
                header_index = min(max(header_row - 1, 0), max(len(raw_rows) - 1, 0))
                if field_names and raw_rows:
                    scan_limit = min(len(raw_rows), max(header_index + 4, 10))

                    def text_cell_count(values: tuple[Any, ...]) -> int:
                        return sum(
                            isinstance(value, str) and bool(value.strip()) for value in values
                        )

                    detected_index = max(
                        range(scan_limit), key=lambda index: text_cell_count(raw_rows[index])
                    )
                    if text_cell_count(raw_rows[detected_index]) >= (
                        text_cell_count(raw_rows[header_index]) + 2
                    ):
                        header_index = detected_index
                raw_headers = raw_rows[header_index] if raw_rows else []
                headers = field_names or [
                    str(value or f"column_{index + 1}")
                    for index, value in enumerate(raw_headers)
                ]
                requested_data_index = (data_start_row - 1) if data_start_row else header_index + 1
                first_data_index = max(requested_data_index, header_index + 1)
                rows = []
                for values in raw_rows[first_data_index:]:
                    if not any(
                        value is not None
                        and (not isinstance(value, str) or bool(value.strip()))
                        for value in values
                    ):
                        continue
                    rows.append(dict(zip(headers, values, strict=False)))
                    if len(rows) >= max_cases:
                        break
            finally:
                workbook.close()
    else:
        raise ValueError(f"Unsupported source format: {extension or '(none)'}")
    normalized = []
    for index, row in enumerate(rows[:max_cases]):
        if isinstance(row, dict):
            normalized.append(json_safe(row))
        else:
            normalized.append({"value": json_safe(row), "_row": index + 1})
    return normalized


def summarize_rows(
    source_format: str,
    rows: list[Any],
    sample_size: int,
    *,
    fields: list[str] | None = None,
    sampled_only: bool = False,
) -> dict[str, Any]:
    samples = [json_safe(row) for row in rows[:sample_size]]
    inferred_fields = fields
    if inferred_fields is None:
        inferred_fields = sorted(
            {str(key) for row in samples if isinstance(row, dict) for key in row}
        )
    return {
        "format": source_format,
        "fields": inferred_fields,
        "sample_count": len(samples),
        "sampled_only": sampled_only,
        "samples": samples,
    }
